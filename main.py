import datetime as dt
from datetime import timedelta
import openpyxl as pyxl
from openpyxl.styles import Font, Border, Side, Alignment
import json
import os

salas = {1:("Alfa", 10), 2:("Beta", 15), 3:("Gamma", 5)} # Tres salas iniciales
clientes = {}
reservaciones = {}
nombresTurnos = {"M": "Matutino", "V": "Vespertino", "N": "Nocturno"}
hoy = dt.datetime.now()

def agregarCliente():
    print("Escribe '0' en cualquier campo para cancelar la operación.")
    while True:
        nombre = input("Ingresa el nombre(s): ")
        if nombre == "" or nombre.strip() == "":
            print("ⓘ El nombre no puede estar vacío.")
            continue
        elif nombre == "0":
            return
        break
    while True:
        apellidos = input("Ingresa los apellidos: ")
        if apellidos == "" or apellidos.strip() == "":
            print("ⓘ Los apellidos no pueden estar vacíos.")
            continue
        elif apellidos == "0":
            return
        break
    if clientes:
        siguienteClave = max(clientes.keys()) + 1
        clientes.update({siguienteClave:[nombre, apellidos]})
    else:
        clientes.update({100:[nombre, apellidos]})
        print("✓ Cliente agregado con éxito.\n")

def mostrarClientes():
    print("\nClientes registrados:")
    print("*"*30)
    print("Clave\tNombre completo")
    for clave, nombreCompleto in clientes.items():
        print(f"{clave}\t{nombreCompleto[0]} {nombreCompleto[1]}")
    print("*"*30)

def reservarSala():
    if not clientes:
        print("ⓘ No hay clientes registrados.")
        return
    while True:
        mostrarClientes()
        try:
            claveCliente = int(input("Ingresa tu clave de cliente: "))
        except ValueError:
            print("⚠︎ Clave inválida.")
            continue
        if claveCliente not in clientes.keys():
            print("⚠︎ La clave de cliente no existe.")
            opcionCancelar = input("¿Cancelar operación? (S/N) ")
            if opcionCancelar.upper() == "S":
                menu()
            elif opcionCancelar.upper() == "N":
                continue
            else:
                print("⚠︎ Opción no reconocida.")
                continue
        else:
            clienteAgendado = claveCliente
            break
    while True:
        fecha_string = input("Ingresa la fecha a agendar (dd/mm/aaaa): ")
        try:
            fechaAgendada = dt.datetime.strptime(fecha_string, "%d/%m/%Y")
        except ValueError:
            print("⚠︎ Fecha inválida.")
            opcionCancelar = input("¿Cancelar operación? (S/N) ")
            if opcionCancelar.upper() == "S":
                return
            elif opcionCancelar.upper() == "N":
                continue
            else:
                print("⚠︎ Opción no reconocida.")
                continue
        if fechaAgendada >= (hoy + timedelta(days=2)):
            break
        else:
            print("ⓘ La reservación tiene que ser hecha con 2 (dos) días de anticipación como mínimo.")
            continue
    todosTurnos = {"M", "V", "N"}
    disponibilidad = {}
    for claveSala in salas.keys():
        disponibilidad.update({claveSala:set(todosTurnos)})

    for llave, (fecha, turno, sala, *_) in reservaciones.items():
        if fecha == fechaAgendada:
            disponibilidad[sala].discard(turno)

    while True:
        print(f"\nSALAS DISPONIBLES Y TURNOS EL {fechaAgendada.strftime("%d %b %Y")}:")
        print("*"*65)
        print(f"{'Clave':<6}\t{'Nombre':<15}\t{'Cupo':<5}\t{'Turnos disponibles':<30}")
        for clave, (nombre, cupo) in salas.items():
            turnos = sorted(disponibilidad[clave])
            turnosLegibles = [nombresTurnos.get(t, t) for t in turnos]
            print(f"{clave:<6}\t{nombre:<15}\t{cupo:<5}\t{", ".join(turnosLegibles):<30}")
        print("*"*65)
        try:
            salaAgendada = int(input("Ingresa la clave de la sala a agendar: "))
        except:
            print("⚠︎ Clave inválida.")
            continue
        if salaAgendada not in salas.keys():
            print("ⓘ La sala no existe.")
            opcionCancelar = input("¿Cancelar operación? (S/N) ")
            if opcionCancelar.upper() == "S":
                return
            elif opcionCancelar.upper() == "N":
                continue
            else:
                print("⚠︎ Opción no reconocida.")
                continue
        if not disponibilidad[salaAgendada]:
            print("ⓘ Esta sala no tiene turnos disponibles en la fecha seleccionada.")
            opcionCancelar = input("¿Cancelar operación? (S/N) ")
            if opcionCancelar.upper() == "S":
                return
            elif opcionCancelar.upper() == "N":
                continue
            else:
                print("⚠︎ Opción no reconocida.")
                continue
        break
    while True:
        turno = input("Elige el turno a agendar (M - matutino, V - vespertino, N - nocturno): ")
        if turno.upper() not in disponibilidad[salaAgendada]:
            print("⚠︎ Turno no disponible para esta sala.")
            opcionCancelar = input("¿Cancelar operación? (S/N) ")
            if opcionCancelar.upper() == "S":
                return
            elif opcionCancelar.upper() == "N":
                continue
            else:
                print("⚠︎ Opción no reconocida.")
                continue
        turnoAgendado = turno.upper()
        break
    while True:
        print("Escribe '0' para cancelar la operación.")
        nombreEvento = input("Ingresa el nombre del evento: ")
        if nombreEvento == "" or nombreEvento.strip() == "":
            print("ⓘ El nombre del evento no puede estar vacío.")
            continue
        break

    if nombreEvento == "0":
        return
    else:
        if reservaciones:
            siguienteClaveEvento = max(reservaciones.keys()) + 1
            reservaciones.update({siguienteClaveEvento:[fechaAgendada, turnoAgendado, salaAgendada, clienteAgendado, nombreEvento]})
        else:
            reservaciones.update({1000:[fechaAgendada, turnoAgendado, salaAgendada, clienteAgendado, nombreEvento]})
        print("✓ La reservación fue registrada con éxito.\n")

def exportarXLSX(fecha:dt.datetime, diccionario:dict):
    fecha_str = dt.datetime.strftime(fecha, '%d %b %Y')
    fecha_str_nums = dt.datetime.strftime(fecha, '%d-%m-%Y')
    wbReporte = pyxl.Workbook()
    hoja = wbReporte["Sheet"]
    hoja.title = "Reservaciones"
    hoja.merge_cells("A1:D1")
    hoja["A1"] = f"RESERVACIONES DEL DIA {fecha_str}"
    encabezados = ["Sala", "Cliente", "Evento", "Turno"]
    hoja.append(encabezados)
    hoja.column_dimensions["A"].width = 10
    hoja.column_dimensions["B"].width = 20
    hoja.column_dimensions["C"].width = 25
    hoja.column_dimensions["D"].width = 10

    for llave, (fechaExportar, turno, salaID, clienteID, nombreEvento) in diccionario.items():
            fila = [salas[salaID][0], clientes[clienteID][0] + ' ' + clientes[clienteID][1], nombreEvento, nombresTurnos[turno]]
            hoja.append(fila)

    negritas = Font(bold = True)
    centrado = Alignment(horizontal="center", vertical="center")
    bordeInferior = Border(bottom=Side(border_style="thick", color="000000"))
    hoja["A1"].font = negritas
    for fila in hoja["A2:D2"]:
        for celda in fila:
            celda.font = negritas
            celda.alignment = centrado
            celda.border = bordeInferior
    for fila in hoja.iter_rows(min_row=1, max_row=hoja.max_row, max_col=hoja.max_column):
        for celda in fila:
            celda.alignment = centrado

    nombreArchivoExportar = f"reporte_{fecha_str_nums}.xlsx"
    wbReporte.save(nombreArchivoExportar)

def consultarReservaciones():
    if not reservaciones:
        print("ⓘ No hay reservaciones. Registra una ahora:")
        reservarSala()
    print("Para consultar una reservación, ingresa la fecha (dd/mm/aaaa) bajo la que fue agendada.")
    print("Escribe '0' en cualquier campo para cancelar la operación.")
    while True:
        fechaConsultada_string = input("Fecha a consultar: ")
        if fechaConsultada_string == "0":
            return
        try:
            fechaConsultada = dt.datetime.strptime(fechaConsultada_string, "%d/%m/%Y")
        except:
            print("⚠︎ Fecha inválida.")
            continue
        consultas = {}
        for llave, valor in reservaciones.items():
            if valor[0] == fechaConsultada:
                consultas.update({llave:valor})

        if not consultas:
            print("ⓘ No existen reservaciones agendadas bajo esta fecha.")
            opcionCancelar = input("¿Cancelar operación? (S/N) ")
            if opcionCancelar.upper() == "S":
                return
            elif opcionCancelar.upper() == "N":
                continue
            else:
                print("⚠︎ Opción no reconocida.")
                continue

        print("\n")
        print("*"*70)
        print(f"REPORTE DE RESERVACIONES PARA EL DIA {fechaConsultada.strftime('%d %b %Y')}")
        print("*"*70)
        print(f"{'Sala':<10}{'Cliente':<20}{'Evento':<30}{'Turno':<15}")
        print("-"*70)
        for claveEvento, detalles in consultas.items():
            fechaAgendada, turnoAgendado, salaAgendada, clienteAgendado, nombreEvento = detalles
            print(f"{salas.get(salaAgendada)[0]:<10}{clientes.get(clienteAgendado)[0] + ' ' + clientes.get(clienteAgendado)[1]:<20}{nombreEvento:<30}{nombresTurnos.get(turnoAgendado):<15}")
        print("*"*70)
        break
    while True:
        opcionExportar = input(f"¿Deseas exportar el reporte de reservaciones del {fechaConsultada.strftime('%d %b %Y')} a Excel? (S/N) ").upper()
        if opcionExportar == "S":
            exportarXLSX(fechaConsultada, consultas)
            print("✓ El reporte fue exportado exitosamente.")
            break
        elif opcionExportar == "N":
            break
        else:
            print("⚠︎ Opción no reconocida.")
            continue

def registrarSala():
    print("Escribe '0' en cualquier campo para cancelar la operación.")
    while True:
        nombreSala = input("Ingresa el nombre de la sala: ")
        if nombreSala == "0":
            return
        if nombreSala == "" or nombreSala.strip() == "":
            print("ⓘ El nombre no puede estar vacío.")
            continue
        break
    while True:
        try:
            cupoSala = int(input("Ingresa el cupo de la sala: "))
        except ValueError:
            print("ⓘ El valor ingresado no es válido. Debe ser un entero.")
            continue
        if cupoSala == 0:
            return
        if salas:
            siguienteSala = max(salas.keys()) + 1
            salas.update({siguienteSala:(nombreSala, cupoSala)})
        print("✓ La sala fue registrada con éxito.\n")
        break

def editarEvento():
    if not reservaciones:
        print("ⓘ No hay reservaciones registradas.")
        return

    while True:
        print("Para editar el nombre de un evento existente, ingresa el rango de fechas (dd/mm/aaaa) en el que se encuentra agendado el evento que quieres editar.")
        print("Escribe '0' en cualquier campo para cancelar la operación.")
        inicioRango_string = input("Del: ")
        if inicioRango_string == "0":
            return
        finRango_string = input("Al: ")
        if finRango_string == "0":
            return
        try:
            inicioRango = dt.datetime.strptime(inicioRango_string, "%d/%m/%Y")
            finRango = dt.datetime.strptime(finRango_string, "%d/%m/%Y")
        except:
            print("⚠︎ Fecha inválida.")
            opcionCancelar = input("¿Cancelar operación? (S/N) ")
            if opcionCancelar.upper() == "S":
                return
            elif opcionCancelar.upper() == "N":
                continue
            else:
                print("⚠︎ Opción no reconocida.")
                continue

        rangoFiltrado = {}
        for llave, valor in reservaciones.items():
            if inicioRango <= valor[0] <= finRango:
                rangoFiltrado.update({llave:valor})

        if not rangoFiltrado:
            print("ⓘ No existen reservaciones en el rango seleccionado.")
            opcionCancelar = input("¿Cancelar operación? (S/N) ")
            if opcionCancelar.upper() == "S":
                return
            elif opcionCancelar.upper() == "N":
                continue
            else:
                print("⚠︎ Opción no reconocida.")
                continue

        print(f"\nEVENTOS REGISTRADOS ENTRE EL {inicioRango.strftime("%d %b %Y")} Y EL {finRango.strftime("%d %b %Y")}:")
        print("*"*100)
        print(f"{'Clave':<10}{'Fecha':<20}{'Nombre':<30}{'Turno':<20}Sala")
        for idEvento, datos in rangoFiltrado.items():
            fecha, turno, sala, cliente, nombreEvento = datos
            print(f"{idEvento:<10}{fecha.strftime('%d/%m/%Y'):<20}{nombreEvento:<30}{nombresTurnos.get(turno):<20}{salas.get(sala)[0]}")
        print("*"*100)
        while True:
            try:
                eventoEditando = int(input("\nIngresa la clave del evento que deseas renombrar: "))
            except ValueError:
                print("⚠︎ Clave inválida.")
                continue
            if eventoEditando == 0:
                return
            if eventoEditando not in rangoFiltrado.keys():
                print("ⓘ Este evento no existe en el rango de fechas.")
                continue
            break
        while True:
            nuevoNombre = input("Ingresa el nuevo nombre para el evento: ")
            if nuevoNombre == "0":
                return
            if nuevoNombre == "" or nuevoNombre.strip() == "":
                print("ⓘ El nombre no puede estar vacío.")
                continue
            break
        reservaciones[eventoEditando][4] = nuevoNombre
        print(f"✓ El nombre del evento con clave {eventoEditando} fue editado a '{nuevoNombre}' exitosamente.\n")

def serializarDatetime(obj):
    if isinstance(obj, dt.datetime):
        return obj.isoformat()
    raise TypeError("Tipo de objeto no serializable")

def guardarJSON():
    if not reservaciones or not clientes:
        print("ⓘ No hay información para guardar.")
        return False

    datosGuardados = {}
    datosGuardados["salas"] = salas
    datosGuardados["clientes"] = clientes
    datosGuardados["reservaciones"] = reservaciones

    with open("estado.json", "w", encoding="utf-8") as archivo:
        json.dump(datosGuardados, archivo, ensure_ascii=False, default=serializarDatetime, indent=4)

    print("✓ Datos del sistema guardados exitosamente en 'estado.json'.")

def revisarJSON():
    if os.path.exists("estado.json") and os.path.getsize("estado.json") != 0:
        with open("estado.json", "r", encoding="utf-8") as archivo:
            datosCargados = json.load(archivo)

        salasCargadas = datosCargados.get("salas", {})
        clientesCargados = datosCargados.get("clientes", {})
        reservacionesCargadas = datosCargados.get("reservaciones", {})

        reservacionesReconstruidas = {}
        for llave, valores in reservacionesCargadas.items():
            copia = valores.copy()
            try:
                copia[0] = dt.datetime.fromisoformat(copia[0])
            except:
                pass
            reservacionesReconstruidas[int(llave)] = copia

        salasReconstruidas = {}
        for llave, valores in salasCargadas.items():
            copia = valores.copy()
            salasReconstruidas[int(llave)] = copia

        clientesReconstruidos = {}
        for llave, valores in clientesCargados.items():
            copia = valores.copy()
            clientesReconstruidos[int(llave)] = copia

        print("ⓘ El archivo 'estado.json' previamente grabado ha sido cargado.")
        return salasReconstruidas, clientesReconstruidos, reservacionesReconstruidas
    else:
        print("ⓘ No hay información guardada previamente. Se iniciará con un archivo en blanco.")
        return salas, {}, {}

def menu():
    while True:
        print("\n")
        print("*"*50)
        print("SISTEMA DE RESERVA DE SALAS PARA COWORKING")
        print("\nSelecciona una opción para continuar:")
        print("(a) Reservar una sala")
        print("(b) Editar el nombre de una reservación")
        print("(c) Consultar reservaciones")
        print("(d) Registrar nuevo cliente")
        print("(e) Registrar nueva sala")
        print("(f) Salir\n")
        print("*"*50)
        opcion = input()
        if opcion.lower() == "a":
            reservarSala()
            continue
        elif opcion.lower() == "b":
            editarEvento()
            continue
        elif opcion.lower() == "c":
            consultarReservaciones()
            continue
        elif opcion.lower() == "d":
            agregarCliente()
            continue
        elif opcion.lower() == "e":
            registrarSala()
            continue
        elif opcion.lower() == "f":
            guardarSalir = input("¿Guardar y salir? (S/N): ")
            if guardarSalir.upper() == "S":
                if guardarJSON() == False:
                    print("Saliendo...")
                    break
                else:
                    print("Saliendo...")
                    break
            elif guardarSalir.upper() == "N":
                continue
        else:
            print("⚠︎ Opción no reconocida.")
            continue

salas, clientes, reservaciones = revisarJSON()
menu()
